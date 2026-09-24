"""Module to handle process finalization"""

import logging
from io import BytesIO

from mbu_msoffice_integration.sharepoint_class import Sharepoint
from mbu_rpa_core.database import RPAConnection
from mbu_rpa_core.exceptions import ProcessError
from openpyxl import load_workbook

from ats_framework.core.application_handler import close, get_app, startup
from ats_framework.core.queue_handler import (
    _clean,
    fetch_all_dagtilbud,
    fetch_submissions,
    parse_dagtilbud,
)
from ats_framework.helpers import config
from ats_framework.helpers.excel_styling import style_report_sheet

logger = logging.getLogger(__name__)

MANGLER_COLUMNS = [
    "Dagtilbud",
    "Adresse",
    "LOSID",
    "Dagtilbudsleder",
    "Leder e-mail",
    "Antal afdelinger",
    "Selvejende",
]


def build_mangler_rows(all_dagtilbud: list[dict], submissions: list[dict]) -> list:
    """Return a row for every dagtilbud that has not submitted the formular."""
    submitted_ids = {
        parse_dagtilbud(
            _clean(s.get("form_data", {}).get("data", {}).get("vaelg_dagtilbud"))
        )["dagtilbud_id"]
        for s in submissions
    }

    return [
        [
            _clean(d.get("ENHNAVN")),
            _clean(d.get("LISADR")),
            _clean(d.get("LOSID")),
            _clean(d.get("LEDERNAVN")),
            _clean(d.get("E_MAIL")),
            d.get("antal_afdelinger"),
            "Ja" if d.get("sdt") == 1 else "Nej",
        ]
        for d in all_dagtilbud
        if _clean(d.get("LOSID")) not in submitted_ids
    ]


def update_mangler_sheet(sharepoint: Sharepoint, rows: list) -> None:
    """Recreate the Mangler sheet in the Excel file, leaving other sheets as is."""
    content = sharepoint.fetch_file_using_open_binary(
        config.EXCEL_FILE_NAME, config.SHAREPOINT_FOLDER_NAME
    )
    if content is None:
        raise ProcessError(f"Could not download '{config.EXCEL_FILE_NAME}'")

    wb = load_workbook(BytesIO(content))

    if config.MANGLER_SHEET_NAME in wb.sheetnames:
        del wb[config.MANGLER_SHEET_NAME]

    ws = wb.create_sheet(config.MANGLER_SHEET_NAME)
    ws.append(MANGLER_COLUMNS)
    for row in rows:
        ws.append(row)

    style_report_sheet(ws)

    stream = BytesIO()
    wb.save(stream)
    sharepoint.upload_file_from_bytes(
        binary_content=stream.getvalue(),
        file_name=config.EXCEL_FILE_NAME,
        folder_name=config.SHAREPOINT_FOLDER_NAME,
    )


def finalize_process():
    """Update the Mangler sheet with dagtilbud that are missing a submission."""
    rpa_conn = RPAConnection(db_env="PROD", commit=False)
    with rpa_conn:
        all_dagtilbud = fetch_all_dagtilbud(rpa_conn)
        submissions = fetch_submissions(rpa_conn)

    rows = build_mangler_rows(all_dagtilbud, submissions)
    logger.info(
        "%d of %d dagtilbud are missing a submission", len(rows), len(all_dagtilbud)
    )

    startup()
    try:
        sharepoint: Sharepoint = get_app()

        files = sharepoint.fetch_files_list(folder_name=config.SHAREPOINT_FOLDER_NAME)
        if files is None:
            raise ProcessError(
                f"Could not list files in SharePoint folder '{config.SHAREPOINT_FOLDER_NAME}'"
            )

        # The file is created when the first submission is processed
        if config.EXCEL_FILE_NAME not in {f["Name"] for f in files}:
            logger.info(
                "'%s' does not exist yet - skipping %s sheet",
                config.EXCEL_FILE_NAME,
                config.MANGLER_SHEET_NAME,
            )
            return

        update_mangler_sheet(sharepoint, rows)

    finally:
        close()
