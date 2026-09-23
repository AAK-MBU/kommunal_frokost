"""Styling of the report sheet, mirroring "Aldersopdelt og type 2024" in helpers"""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

FONT_NAME = "Arial"

# Header colours per column group (theme colours from the example resolved to RGB)
MASTERDATA_FILL = "C4BD97"
BESVARELSE_FILL = "E6B9B8"
META_FILL = "B7DEE8"
DAGTILBUD_DATA_FILL = "8EB4E3"

# Column name -> (header fill, column width)
COLUMN_STYLES = {
    "Dagtilbud": (MASTERDATA_FILL, 18),
    "Dagtilbudsleder": (MASTERDATA_FILL, 22),
    "LISID": (MASTERDATA_FILL, 8),
    "Enhedens kaldenavn": (MASTERDATA_FILL, 28),
    "Afdelingstype": (MASTERDATA_FILL, 17),
    "Ledertype": (MASTERDATA_FILL, 20),
    "Navn på leder": (MASTERDATA_FILL, 22),
    "Leder e-mail": (MASTERDATA_FILL, 26),
    "Ejertype": (MASTERDATA_FILL, 11),
    "Bestyrelsens beslutning (frokost/madpakke)": (BESVARELSE_FILL, 20),
    "Aldersopdelt valg ja/nej": (BESVARELSE_FILL, 14),
    "Type måltid (egenproduktion/samproduktion eller ekstern)": (BESVARELSE_FILL, 26),
    "Hvorfra samproduktion": (BESVARELSE_FILL, 24),
    "Ekstern leverandør": (BESVARELSE_FILL, 24),
    "Bemærkning": (BESVARELSE_FILL, 34),
    "Indsendt": (META_FILL, 16),
    "Besvarelses-ID": (META_FILL, 13),
}

HEADER_ROW_HEIGHT = 38.25

THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _fill(rgb: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=rgb, end_color=rgb)


def style_report_sheet(ws: Worksheet) -> None:
    """
    Apply the report styling to a sheet with a single header row.

    Has to run after format_and_sort_excel_file, since that rewrites all rows
    and resets fonts and alignment.
    """
    headers = [cell.value for cell in ws[1]]

    for col_idx, header in enumerate(headers, start=1):
        header_fill, width = COLUMN_STYLES.get(header, (MASTERDATA_FILL, 15))
        is_dagtilbud = header == "Dagtilbud"

        column = ws.iter_rows(min_col=col_idx, max_col=col_idx)
        for (cell,) in column:
            cell.border = BORDER

            if cell.row == 1:
                cell.fill = _fill(header_fill)
                cell.font = Font(name=FONT_NAME, size=9, bold=True)
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            else:
                cell.font = Font(name=FONT_NAME, size=8, bold=is_dagtilbud)
                cell.alignment = Alignment(
                    horizontal="left", vertical="top", wrap_text=True
                )
                if is_dagtilbud:
                    cell.fill = _fill(DAGTILBUD_DATA_FILL)

        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    # Let Excel size data rows to the wrapped content
    for row_idx in list(ws.row_dimensions):
        ws.row_dimensions[row_idx].height = None
    ws.row_dimensions[1].height = HEADER_ROW_HEIGHT

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
